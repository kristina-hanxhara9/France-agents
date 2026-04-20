"""
Retailer enrichment CLI — skeleton.

Commands:
  prepare --limit N   Write up to N un-enriched rows to web_enrich_queue.json.
                     Prints "N companies to search" or "0 companies to search".
  save --id ID --json '{...}'
                     Persist one enriched record to web_enrich_results.json.
  status             Print progress counts.
  compile            Merge results into france_retailers-with-keywords.enriched.xlsx.

State files:
  web_enrich_queue.json     Current batch handed to the agent.
  web_enrich_results.json   All saved results, keyed by id. Survives across runs.

Expected xlsx columns (case-insensitive, auto-detected):
  id, company_name, trade_name, search_query
Any of: keywords, siret, city, ...  are carried through but not required.

Adjust COLUMN_MAP below to match your actual xlsx headers.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook, Workbook

ROOT = Path(__file__).parent
XLSX_PATH = ROOT / "france_retailers-with-keywords.xlsx"
OUT_XLSX = ROOT / "france_retailers-with-keywords.enriched.xlsx"
QUEUE_PATH = ROOT / "web_enrich_queue.json"
RESULTS_PATH = ROOT / "web_enrich_results.json"

# Adjust to your real header names. Left side = canonical key, right side = header.
COLUMN_MAP = {
    "id": "id",
    "company_name": "company_name",
    "trade_name": "trade_name",
    "search_query": "search_query",
}

ENRICH_FIELDS = [
    "website",
    "phone",
    "email",
    "web_description",
    "web_products",
    "web_business_type",
    "web_channel_guess",
    "web_channel_detail",
    "facebook",
    "instagram",
    "linkedin",
    "twitter",
]


def _read_rows() -> list[dict]:
    if not XLSX_PATH.exists():
        sys.exit(f"ERROR: {XLSX_PATH} not found")
    wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    lower_headers = [h.lower() for h in headers]

    def col(name: str) -> int | None:
        target = COLUMN_MAP.get(name, name).lower()
        return lower_headers.index(target) if target in lower_headers else None

    idx = {k: col(k) for k in COLUMN_MAP}
    out: list[dict] = []
    for i, row in enumerate(rows[1:], start=2):
        rec = {"_excel_row": i}
        for key, col_idx in idx.items():
            rec[key] = (row[col_idx] if col_idx is not None and col_idx < len(row) else "") or ""
        # Keep every header for pass-through on compile.
        rec["_all"] = {headers[j]: (row[j] if j < len(row) else "") for j in range(len(headers))}
        # Auto-build search_query if missing.
        if not rec.get("search_query"):
            parts = [str(rec.get("company_name", "")).strip()]
            tn = str(rec.get("trade_name", "")).strip()
            if tn and tn.lower() != parts[0].lower():
                parts.append(tn)
            rec["search_query"] = " ".join(parts + ["france"]).strip()
        # Fabricate id if missing.
        if not rec.get("id"):
            rec["id"] = str(i)
        else:
            rec["id"] = str(rec["id"])
        out.append(rec)
    return out


def _load_results() -> dict:
    if RESULTS_PATH.exists():
        return json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    return {}


def _save_results(data: dict) -> None:
    RESULTS_PATH.write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def cmd_prepare(limit: int) -> None:
    rows = _read_rows()
    done = _load_results()
    pending = [r for r in rows if r["id"] not in done]
    batch = pending[:limit]
    queue = [
        {
            "id": r["id"],
            "company_name": r.get("company_name", ""),
            "trade_name": r.get("trade_name", ""),
            "search_query": r.get("search_query", ""),
        }
        for r in batch
    ]
    QUEUE_PATH.write_text(
        json.dumps(queue, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if not queue:
        print("0 companies to search")
    else:
        print(f"{len(queue)} companies to search ({len(pending) - len(queue)} more after this batch)")


def cmd_save(company_id: str, payload: str) -> None:
    try:
        data = json.loads(payload)
    except json.JSONDecodeError as e:
        sys.exit(f"ERROR: invalid JSON: {e}")
    if not isinstance(data, dict):
        sys.exit("ERROR: --json must be an object")
    cleaned = {f: str(data.get(f, "") or "") for f in ENRICH_FIELDS}
    results = _load_results()
    results[str(company_id)] = cleaned
    _save_results(results)
    print(f"saved id={company_id}")


def cmd_status() -> None:
    rows = _read_rows()
    done = _load_results()
    total = len(rows)
    enriched = len(done)
    with_web = sum(1 for v in done.values() if v.get("website"))
    with_phone = sum(1 for v in done.values() if v.get("phone"))
    with_email = sum(1 for v in done.values() if v.get("email"))
    print(f"total rows:        {total}")
    print(f"enriched:          {enriched}")
    print(f"remaining:         {total - enriched}")
    print(f"  with website:    {with_web}")
    print(f"  with phone:      {with_phone}")
    print(f"  with email:      {with_email}")


def cmd_compile() -> None:
    rows = _read_rows()
    results = _load_results()
    if not rows:
        sys.exit("ERROR: no rows to compile")

    headers = list(rows[0]["_all"].keys())
    for f in ENRICH_FIELDS:
        if f not in headers:
            headers.append(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "retailers"
    ws.append(headers)
    for r in rows:
        enriched = results.get(r["id"], {})
        merged = dict(r["_all"])
        for f in ENRICH_FIELDS:
            merged[f] = enriched.get(f, merged.get(f, ""))
        ws.append([merged.get(h, "") for h in headers])
    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX} ({len(rows)} rows, {len(results)} enriched)")


def main() -> None:
    parser = argparse.ArgumentParser(prog="retailer_enrich")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_prep = sub.add_parser("prepare")
    p_prep.add_argument("--limit", type=int, default=10)

    p_save = sub.add_parser("save")
    p_save.add_argument("--id", required=True)
    p_save.add_argument("--json", required=True)

    sub.add_parser("status")
    sub.add_parser("compile")

    args = parser.parse_args()
    if args.cmd == "prepare":
        cmd_prepare(args.limit)
    elif args.cmd == "save":
        cmd_save(args.id, args.json)
    elif args.cmd == "status":
        cmd_status()
    elif args.cmd == "compile":
        cmd_compile()


if __name__ == "__main__":
    main()
